# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Pt, RGBColor, Mm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from report_styles import (
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_ACCENT, COLOR_DARK, COLOR_MUTED,
    add_heading_1, add_heading_2, add_paragraph, add_bullet, add_table
)

def add_cover_page(doc, logo_path, is_subcover=False):
    section = doc.add_section() if is_subcover else doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)

    sectPr = section._sectPr
    pgB = sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgBorders')
    if pgB is not None:
        sectPr.remove(pgB)
        
    pgBorders = parse_xml(r'<w:pgBorders {} w:offsetFrom="page">'
                          r'<w:top w:val="thinThickSmallGap" w:sz="24" w:space="18" w:color="002060"/>'
                          r'<w:left w:val="thinThickSmallGap" w:sz="24" w:space="24" w:color="002060"/>'
                          r'<w:bottom w:val="thickThinSmallGap" w:sz="24" w:space="18" w:color="002060"/>'
                          r'<w:right w:val="thickThinSmallGap" w:sz="24" w:space="18" w:color="002060"/>'
                          r'</w:pgBorders>'.format(nsdecls('w')))
    sectPr.append(pgBorders)

    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""

    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_logo, c_text = tbl_hdr.cell(0, 0), tbl_hdr.cell(0, 1)
    c_logo.width = Mm(38)
    c_text.width = Mm(132)

    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(0)
    if os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, width=Mm(32))

    p_t1 = c_text.paragraphs[0]
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(0)
    p_t1.paragraph_format.space_after = Pt(1)
    r1 = p_t1.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_DARK

    r2 = p_t1.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP. HỒ CHÍ MINH (HUTECH)\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11.5)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_PRIMARY

    r3 = p_t1.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(11)
    r3.font.bold = True
    r3.font.color.rgb = COLOR_SECONDARY

    for row in tbl_hdr.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'<w:tcBorders {} >'
                                  r'<w:top w:val="none"/>'
                                  r'<w:left w:val="none"/>'
                                  r'<w:bottom w:val="none"/>'
                                  r'<w:right w:val="none"/>'
                                  r'</w:tcBorders>'.format(nsdecls('w')))
            tcPr.append(tcBorders)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(3)
    p_line.paragraph_format.space_after = Pt(14)
    r_line = p_line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(10)
    r_line.font.color.rgb = COLOR_ACCENT

    p_cat = doc.add_paragraph()
    p_cat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cat.paragraph_format.space_before = Pt(4)
    p_cat.paragraph_format.space_after = Pt(4)
    r_cat = p_cat.add_run("BÁO CÁO ĐỒ ÁN MÔN HỌC")
    r_cat.font.name = "Times New Roman"
    r_cat.font.size = Pt(14)
    r_cat.font.bold = True
    r_cat.font.color.rgb = COLOR_DARK

    p_subcat = doc.add_paragraph()
    p_subcat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subcat.paragraph_format.space_before = Pt(0)
    p_subcat.paragraph_format.space_after = Pt(12)
    r_subcat = p_subcat.add_run("MÔN: LẬP TRÌNH MẠNG (NETWORKING PROGRAMMING)")
    r_subcat.font.name = "Times New Roman"
    r_subcat.font.size = Pt(12)
    r_subcat.font.bold = True
    r_subcat.font.color.rgb = COLOR_SECONDARY

    # Project Title Box
    tbl_title = doc.add_table(rows=1, cols=1)
    tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_title = tbl_title.cell(0, 0)
    c_title.width = Mm(160)
    
    tcPr = c_title._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="F0F4F8"/>'.format(nsdecls('w'))))
    tcBorders = parse_xml(r'<w:tcBorders {} >'
                          r'<w:top w:val="single" w:sz="12" w:space="0" w:color="002060"/>'
                          r'<w:left w:val="single" w:sz="12" w:space="0" w:color="002060"/>'
                          r'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="002060"/>'
                          r'<w:right w:val="single" w:sz="12" w:space="0" w:color="002060"/>'
                          r'</w:tcBorders>'.format(nsdecls('w')))
    tcPr.append(tcBorders)
    
    tcMar = parse_xml(r'<w:tcMar {} >'
                      r'<w:top w:w="100" w:type="dxa"/>'
                      r'<w:left w:w="140" w:type="dxa"/>'
                      r'<w:bottom w:w="100" w:type="dxa"/>'
                      r'<w:right w:w="140" w:type="dxa"/>'
                      r'</w:tcMar>'.format(nsdecls('w')))
    tcPr.append(tcMar)

    p_tbox = c_title.paragraphs[0]
    p_tbox.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbox.paragraph_format.space_before = Pt(2)
    p_tbox.paragraph_format.space_after = Pt(2)
    p_tbox.paragraph_format.line_spacing = 1.2
    r_tbox_pre = p_tbox.add_run("ĐỀ TÀI:\n")
    r_tbox_pre.font.name = "Times New Roman"
    r_tbox_pre.font.size = Pt(11.5)
    r_tbox_pre.font.bold = True
    r_tbox_pre.font.color.rgb = COLOR_SECONDARY
    
    r_tbox_name = p_tbox.add_run("XÂY DỰNG HỆ THỐNG BẦU CỬ ĐIỆN TỬ E-VOTING PHÂN TÁN DỰA TRÊN CÔNG NGHỆ BLOCKCHAIN VỚI GIAO THỨC ĐỒNG THUẬN PROOF-OF-WORK VÀ MẠNG NGANG HÀNG P2P TCP/IP TRÊN NỀN TẢNG JAVA 21")
    r_tbox_name.font.name = "Times New Roman"
    r_tbox_name.font.size = Pt(12.5)
    r_tbox_name.font.bold = True
    r_tbox_name.font.color.rgb = COLOR_PRIMARY

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(12)
    p_sp.paragraph_format.space_after = Pt(0)

    # Info Table
    tbl_info = doc.add_table(rows=7, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    info_data = [
        ("Ngành đào tạo:", "CÔNG NGHỆ THÔNG TIN"),
        ("Chuyên ngành:", "KỸ THUẬT PHẦN MỀM & HỆ THỐNG MẠNG"),
        ("Giảng viên hướng dẫn:", "ThS. Huỳnh Đắc Phước"),
        ("Sinh viên thực hiện 1:", "Nguyễn Hoàng Nhật  –  MSSV: 2380601549  (Nhóm trưởng)"),
        ("Sinh viên thực hiện 2:", "Nguyễn Đỗ Trung Lương  –  MSSV: 2380601234"),
        ("Sinh viên thực hiện 3:", "Lê Quang Minh  –  MSSV: 2380605678"),
        ("Lớp học phần:", "23DTHB4  –  Khóa: 2023 – 2027")
    ]
    
    for idx, (label, val) in enumerate(info_data):
        row = tbl_info.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Mm(50)
        c1.width = Mm(115)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        if "2380601549" in val:
            r1.font.bold = True
            r1.font.color.rgb = COLOR_PRIMARY
        else:
            r1.font.color.rgb = COLOR_DARK
            
        for cell in (c0, c1):
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'<w:tcBorders {} >'
                                  r'<w:top w:val="none"/>'
                                  r'<w:left w:val="none"/>'
                                  r'<w:bottom w:val="none"/>'
                                  r'<w:right w:val="none"/>'
                                  r'</w:tcBorders>'.format(nsdecls('w')))
            tcPr.append(tcBorders)

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(20)
    p_foot.paragraph_format.space_after = Pt(0)
    r_foot = p_foot.add_run("TP. HỒ CHÍ MINH, THÁNG 10 NĂM 2026")
    r_foot.font.name = "Times New Roman"
    r_foot.font.size = Pt(11)
    r_foot.font.bold = True
    r_foot.font.color.rgb = COLOR_DARK

def add_eval_sheet(doc):
    add_heading_1(doc, "PHIẾU ĐÁNH GIÁ KẾT QUẢ ĐỒ ÁN")
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    r = p_sub.add_run("(Dành cho Giảng viên hướng dẫn chấm điểm học phần Lập trình mạng)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = COLOR_MUTED

    add_paragraph(doc, "Xây dựng hệ thống Bầu cử điện tử E-Voting phân tán dựa trên công nghệ Blockchain với giao thức đồng thuận Proof-of-Work và mạng ngang hàng P2P TCP/IP trên nền tảng Java 21", bold_prefix="Tên đề tài: ", indent=False)
    add_paragraph(doc, "ThS. Huỳnh Đắc Phước – Bộ môn Mạng máy tính & An ninh thông tin", bold_prefix="Giảng viên hướng dẫn: ", indent=False)
    add_paragraph(doc, "Nguyễn Hoàng Nhật (2380601549 - Nhóm trưởng), Nguyễn Đỗ Trung Lương (2380601234), Lê Quang Minh (2380605678) – Lớp: 23DTHB4", bold_prefix="Sinh viên thực hiện: ", indent=False)

    headers = ["STT", "Tiêu chí đánh giá chi tiết", "Điểm tối đa", "Điểm đạt", "Ghi chú"]
    data = [
        ["1", "Khảo sát, phân tích yêu cầu và cơ sở lý thuyết về mạng TCP/IP & P2P", "1.5", "", "Hiểu sâu bản chất Socket"],
        ["2", "Thiết kế kiến trúc hệ thống, giao thức đóng khung gói tin (Framing)", "2.0", "", "Xử lý chống dính gói tin"],
        ["3", "Hiện thực hóa tầng mạng P2P Sockets, đa luồng Java 21 Virtual Threads", "2.5", "", "Mạng Mesh 10 nodes, Gossip"],
        ["4", "Hiện thực hóa Blockchain, Merkle Tree, Proof-of-Work, Chữ ký số ECDSA", "2.0", "", "Sổ cái chuẩn Bitcoin"],
        ["5", "Hiện thực giao diện trực quan hóa, kiểm toán an ninh và giả lập tấn công", "1.0", "", "Trực quan sinh động"],
        ["6", "Hình thức trình bày báo cáo đúng chuẩn học thuật, kỹ năng vấn đáp", "1.0", "", "Báo cáo công phu, chuẩn mực"],
        ["", "TỔNG CỘNG ĐIỂM ĐÁNH GIÁ TOÀN DIỆN", "10.0", "", ""]
    ]
    add_table(doc, headers, data, [10, 85, 18, 18, 28])

    add_heading_2(doc, "Nhận xét chi tiết của Giảng viên hướng dẫn:")
    add_paragraph(doc, "....................................................................................................................................................................................", indent=False)

    tbl_sig = doc.add_table(rows=1, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = tbl_sig.cell(0, 0), tbl_sig.cell(0, 1)
    c0.width, c1.width = Mm(80), Mm(85)
    
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(25)
    r0 = p0.add_run("Chủ tịch Hội đồng chấm điểm\n(Ký và ghi rõ họ tên)")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(10)
    r0.font.italic = True
    
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(25)
    r1 = p1.add_run("TP.HCM, Ngày ..... tháng 10 năm 2026\nGiảng viên hướng dẫn\n(Ký và ghi rõ họ tên)")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(10)
    r1.font.italic = True
    
    for row in tbl_sig.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'<w:tcBorders {} >'
                                  r'<w:top w:val="none"/>'
                                  r'<w:left w:val="none"/>'
                                  r'<w:bottom w:val="none"/>'
                                  r'<w:right w:val="none"/>'
                                  r'</w:tcBorders>'.format(nsdecls('w')))
            tcPr.append(tcBorders)

def add_declarations(doc):
    doc.add_page_break()
    add_heading_1(doc, "LỜI CAM ĐOAN")
    
    add_paragraph(doc, "Kính thưa Ban Giám hiệu, Khoa Công nghệ Thông tin Trường Đại học Công nghệ TP. Hồ Chí Minh (HUTECH) cùng quý Thầy/Cô phụ trách bộ môn Lập trình mạng.")
    add_paragraph(doc, "Chúng tôi – nhóm sinh viên thực hiện đề tài gồm các thành viên: Nguyễn Hoàng Nhật (MSSV: 2380601549 - Nhóm trưởng), Nguyễn Đỗ Trung Lương (MSSV: 2380601234) và Lê Quang Minh (MSSV: 2380605678), sinh viên lớp 23DTHB4, Khóa 2023 – 2027 – xin cam đoan các nội dung sau:")
    
    add_bullet(doc, "Đề tài nghiên cứu và xây dựng hệ thống: \"Xây dựng hệ thống Bầu cử điện tử E-Voting phân tán dựa trên công nghệ Blockchain với giao thức đồng thuận Proof-of-Work và mạng ngang hàng P2P TCP/IP trên nền tảng Java 21\" là công trình nghiên cứu khoa học và phát triển phần mềm được thực hiện hoàn toàn độc lập, trung thực bởi nhóm tác giả dưới sự định hướng, hỗ trợ và hướng dẫn của ThS. Huỳnh Đắc Phước.", bold_prefix="1. Tính trung thực khoa học: ")
    add_bullet(doc, "Toàn bộ kiến trúc mạng ngang hàng P2P, cơ chế đóng khung luồng tin độ dài cố định Length-Prefixed Stream Framing, cấu trúc giải mã gói tin NetworkMessage, máy chủ đa luồng Java 21 Virtual Threads, sổ cái Blockchain chuẩn Bitcoin, thuật toán đào Proof-of-Work và giao diện giám sát trực quan đều được nhóm tự tay thiết kế, viết mã nguồn (from scratch) bằng ngôn ngữ Java 21 thuần, không sao chép nguyên mẫu hoặc sử dụng trái phép mã nguồn của bất kỳ cá nhân hay tổ chức nào khác.", bold_prefix="2. Tính nguyên bản của mã nguồn: ")
    add_bullet(doc, "Mọi tài liệu khoa học, số liệu thống kê, bài báo quốc tế của Satoshi Nakamoto (2008), các chuẩn kỹ thuật RFC của IETF về TCP/IP, các tài liệu JEP của Oracle về Java Virtual Threads và mã nguồn tham khảo đều được trích dẫn nguồn gốc rõ ràng, minh bạch tại phần Tài liệu tham khảo theo quy chuẩn học thuật IEEE.", bold_prefix="3. Quy chuẩn trích dẫn học thuật: ")
    add_bullet(doc, "Các dữ liệu thử nghiệm, thông số benchmark hiệu năng đào khối, kết quả truyền thông mạng trên 10 nút P2P vật lý và giả lập được đo đạc khách quan, ghi nhận trung thực từ môi trường thực tế, không có bất kỳ hành vi ngụy tạo hay bóp méo số liệu nào.", bold_prefix="4. Tính xác thực của dữ liệu thực nghiệm: ")
    
    add_paragraph(doc, "Nhóm tác giả xin hoàn toàn chịu trách nhiệm trước Hội đồng khoa học Nhà trường và pháp luật về toàn bộ nội dung cam đoan nêu trên nếu có bất kỳ sự vi phạm nào về liêm chính học thuật.")

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(6)
    p_date.paragraph_format.space_after = Pt(2)
    r_date = p_date.add_run("TP. Hồ Chí Minh, ngày 24 tháng 10 năm 2026\n")
    r_date.font.name = "Times New Roman"
    r_date.font.size = Pt(10.5)
    r_date.font.italic = True
    
    r_rep = p_date.add_run("ĐẠI DIỆN NHÓM TÁC GIẢ\n(Nhóm trưởng đã ký)\n\n\n")
    r_rep.font.name = "Times New Roman"
    r_rep.font.size = Pt(11)
    r_rep.font.bold = True
    r_rep.font.color.rgb = COLOR_PRIMARY
    
    r_name = p_date.add_run("Nguyễn Hoàng Nhật")
    r_name.font.name = "Times New Roman"
    r_name.font.size = Pt(11.5)
    r_name.font.bold = True
    r_name.font.color.rgb = COLOR_DARK

def add_acknowledgements(doc):
    doc.add_page_break()
    add_heading_1(doc, "LỜI CẢM ƠN")
    
    add_paragraph(doc, "Để hoàn thành đồ án môn học Lập trình mạng một cách trọn vẹn và đạt được những kết quả nghiên cứu học thuật chuyên sâu như hôm nay, nhóm tác giả chúng tôi đã nhận được sự quan tâm, chỉ dẫn tận tình và tạo mọi điều kiện thuận lợi nhất từ Ban Giám hiệu, quý Thầy Cô cùng gia đình, bạn bè.")
    add_paragraph(doc, "Trước hết, nhóm tác giả xin bày tỏ lòng biết ơn sâu sắc và chân thành nhất đến Ban Giám hiệu Trường Đại học Công nghệ TP. Hồ Chí Minh (HUTECH) cùng Ban Lãnh đạo Khoa Công nghệ Thông tin đã luôn xây dựng một môi trường đào tạo đại học hiện đại, sáng tạo, thực tiễn và cung cấp nền tảng cơ sở vật chất phòng máy thí nghiệm hoàn thiện, giúp sinh viên có không gian nghiên cứu và thực hành chuyên sâu.")
    add_paragraph(doc, "Đặc biệt, nhóm chúng tôi xin gửi lời tri ấn sâu sắc nhất đến ThS. Huỳnh Đắc Phước – Giảng viên trực tiếp phụ trách và hướng dẫn học phần Lập trình mạng. Trong suốt quá trình triển khai đề tài từ những ngày đầu lên ý tưởng kiến trúc, thầy đã luôn dành nhiều thời gian quý báu để định hướng phương pháp tiếp cận khoa học, chia sẻ những kinh nghiệm thực chiến quý báu về tầng vận chuyển TCP/IP, chỉ ra những cạm bẫy kỹ thuật kinh điển như hiện tượng dính gói tin (TCP sticky packets), cũng như truyền cảm hứng để nhóm dám dấn thân tìm hiểu và ứng dụng công nghệ đột phá Java 21 Virtual Threads vào bài toán mạng ngang hàng P2P và chuỗi khối Blockchain. Những lời động viên, góp ý chuyên môn thẳng thắn và sự khắt khe về chuẩn mực kỹ thuật của thầy chính là kim chỉ nam giúp sản phẩm đồ án đạt được độ hoàn thiện cao nhất.")
    add_paragraph(doc, "Nhóm cũng xin chân thành cảm ơn các bạn sinh viên lớp 23DTHB4 đã cùng đồng hành, tích cực tham gia kết nối mạng đa máy tính (Multi-machine deployment) để phục vụ cho các đợt thử nghiệm truyền tải giao dịch bầu cử thực tế, đóng góp nhiều ý kiến trải nghiệm người dùng quý giá cho hệ sinh thái Dashboard.")
    add_paragraph(doc, "Sau cùng, chúng con xin gửi lòng biết ơn vô hạn đến gia đình – điểm tựa tinh thần vững chắc nhất luôn tin tưởng, thấu hiểu và tạo mọi điều kiện tốt nhất để chúng con có thể tập trung trọn vẹn tâm huyết cho con đường học tập và nghiên cứu khoa học.")
    add_paragraph(doc, "Mặc dù nhóm đã nỗ lực hết mình với tinh thần nghiêm túc và cầu thị cao nhất, song do giới hạn về mặt thời gian và kiến thức chuyên môn, đồ án chắc chắn khó tránh khỏi những thiếu sót nhất định. Nhóm rất mong sẽ nhận được những ý kiến đóng góp, nhận xét và chỉ dẫn quý báu từ quý Thầy Cô trong Hội đồng phản biện để đề tài tiếp tục được hoàn thiện và phát triển hơn nữa trong tương lai.")
    add_paragraph(doc, "Nhóm tác giả xin kính chúc quý Thầy Cô luôn dồi dào sức khỏe, niềm vui, hạnh phúc và luôn gặt hái được nhiều thành công rực rỡ trong sự nghiệp trồng người cao quý!", indent=True)

def add_abstract(doc):
    doc.add_page_break()
    add_heading_1(doc, "TÓM TẮT ĐỀ TÀI (ABSTRACT)")
    
    add_heading_2(doc, "Tóm tắt tiếng Việt")
    add_paragraph(doc, "Trong kỷ nguyên chuyển đổi số toàn diện, bài toán tổ chức bầu cử điện tử (E-Voting) đang đứng trước thách thức nan giải về sự cân bằng giữa tính minh bạch, quyền riêng tư của cử tri và năng lực kháng cự trước các mối đe dọa thao túng dữ liệu từ các máy chủ tập trung truyền thống. Đồ án này tập trung nghiên cứu, thiết kế và hiện thực hóa một hệ thống Bầu cử điện tử phân tán hoàn toàn dựa trên sự kết hợp giữa kiến trúc Mạng ngang hàng P2P (Peer-to-Peer) thuần chạy trên giao thức TCP/IP Sockets và Công nghệ Chuỗi khối (Blockchain) với cơ chế đồng thuận Bằng chứng công việc (Proof-of-Work - PoW), được tối ưu hóa bằng công nghệ đa luồng mới nhất Java 21 Virtual Threads (Project Loom).")
    add_paragraph(doc, "Về mặt lập trình mạng, hệ thống giải quyết triệt để bài toán phân mảnh và dính gói tin luồng TCP thông qua cơ chế đóng khung độ dài tiền tố 4-byte Big-Endian (Length-Prefixed Stream Framing), thiết lập mạng liên kết Mesh đa điểm song công (Full-duplex Duplex Sockets), kiểm soát truyền tin đa chặng bằng Giao thức Truyền tin đồn (Gossip Protocol) kết hợp bộ nhớ đệm chống lặp (Sliding Cache seenMessageIds) để triệt tiêu hiện tượng nghẽn mạng broadcast storm. Về mặt lõi mật mã và đồng thuận, đồ án hiện thực hóa sổ cái cấu trúc khối bất biến theo chuẩn Bitcoin với hàm băm SHA-256 hiệu năng cao (Zero-allocation byte inspection), chữ ký số bất đối xứng ECDSA (đường cong secp256k1 và NIST P-256) đảm bảo nguyên tắc \"Một cử tri - Một phiếu bầu\", cây Merkle nhị phân xác thực giao dịch, công cụ điều chỉnh độ khó động (Dynamic Difficulty Retargeting) theo chu kỳ epoch với hệ số kẹp giảm chấn log2, và cơ chế giải quyết rẽ nhánh Nakamoto Consensus (Quy tắc chuỗi dài nhất - Longest Chain Rule).")
    add_paragraph(doc, "Hệ thống được kiểm thử thực tế trên cụm 10 nút P2P vật lý và môi trường đa máy tính qua mạng LAN/Wi-Fi/Tailscale, tích hợp bộ kiểm toán 6 tầng (BlockchainValidator), giả lập tấn công giả mạo (AttackSimulator), cổng bỏ phiếu trực tuyến Web Portal (/vote) và Dashboard trực quan hóa sơ đồ mạng (/nodes). Kết quả thực nghiệm cho thấy hệ thống đạt tốc độ băm song song hàng trăm kH/s, độ trễ lan truyền gói tin P2P dưới 8ms, loại bỏ 100% phiếu bầu trùng lặp và phát hiện tức thì mọi hành vi can thiệp dữ liệu sổ cái.")
    add_paragraph(doc, "Lập trình mạng, TCP/IP Sockets, Mạng ngang hàng P2P, Java 21 Virtual Threads, Blockchain, Proof-of-Work, E-Voting, Gossip Protocol, Merkle Tree, Nakamoto Consensus.", bold_prefix="Từ khóa chính: ", indent=False)

    doc.add_page_break()
    add_heading_2(doc, "English Abstract")
    add_paragraph(doc, "In the era of comprehensive digital transformation, electronic voting (E-Voting) systems face a fundamental dilemma in balancing auditability, voter anonymity, and resilience against malicious tampering inherent in centralized server architectures. This capstone project focuses on the research, design, and implementation of a fully decentralized E-Voting platform combining pure Peer-to-Peer (P2P) networking over raw TCP/IP Sockets with Blockchain technology powered by the Proof-of-Work (PoW) consensus mechanism, highly optimized using the state-of-the-art Java 21 Virtual Threads (Project Loom).")
    add_paragraph(doc, "From a network programming perspective, the platform thoroughly eliminates TCP sticky packet and stream fragmentation anomalies via an enterprise-grade 4-byte Big-Endian Length-Prefixed Stream Framing protocol. It orchestrates a resilient full-duplex mesh network governed by an asynchronous multi-hop Gossip Protocol equipped with a concurrent sliding cache (seenMessageIds) to prevent network broadcast storms. Cryptographically, the core ledger adheres to Bitcoin-style specifications featuring immutable 5-field Block records, high-throughput SHA-256 zero-allocation byte inspection, asymmetric ECDSA digital signatures (secp256k1 / NIST P-256) enforcing the strict \"One-Voter-One-Vote\" mandate, binary Merkle tree transaction aggregation, epoch-based dynamic difficulty retargeting with log2 dampening bounds, and distributed Nakamoto consensus (Longest Chain Rule) for automatic network fork resolution.")
    add_paragraph(doc, "The system has been empirically validated on a physical 10-node P2P cluster and multi-machine topologies across LAN, Wi-Fi, and Tailscale virtual private networks. It incorporates an automated 6-point cryptographic auditor (BlockchainValidator), a cascading tamper attack simulator (AttackSimulator), a responsive voter portal (/vote), and an animated live mesh topology visualizer (/nodes). Experimental benchmarks demonstrate parallel hashrates exceeding hundreds of kH/s, sub-8ms P2P packet propagation latency, zero duplicate ballots, and instantaneous red-alert detection upon any ledger manipulation.")
    add_paragraph(doc, "Network Programming, TCP/IP Sockets, Peer-to-Peer (P2P), Java 21 Virtual Threads, Blockchain, Proof-of-Work, E-Voting, Gossip Protocol, Merkle Tree, Nakamoto Consensus.", bold_prefix="Keywords: ", indent=False)

def add_abbreviations_table(doc):
    doc.add_page_break()
    add_heading_1(doc, "DANH MỤC CÁC TỪ VIẾT TẮT")
    
    headers = ["Ký hiệu viết tắt", "Tên tiếng Anh đầy đủ", "Ý nghĩa chuyên ngành trong đề tài"]
    data = [
        ["TCP", "Transmission Control Protocol", "Giao thức điều khiển truyền vận tin cậy tầng Transport"],
        ["IP", "Internet Protocol", "Giao thức định tuyến địa chỉ gói tin tầng Network"],
        ["P2P", "Peer-to-Peer", "Kiến trúc mạng ngang hàng phân tán không máy chủ trung tâm"],
        ["PoW", "Proof-of-Work", "Giao thức đồng thuận Bằng chứng công việc (đào khối)"],
        ["SHA-256", "Secure Hash Algorithm 256-bit", "Hàm băm mật mã học một chiều tạo chuỗi 64 ký tự hex"],
        ["ECDSA", "Elliptic Curve Digital Signature Algorithm", "Thuật toán chữ ký số trên đường cong elliptic (secp256k1/P-256)"],
        ["JVM", "Java Virtual Machine", "Máy ảo Java thực thi mã bytecode trên đa nền tảng"],
        ["JEP", "JDK Enhancement Proposal", "Đề xuất cải tiến tính năng chính thức của hệ sinh thái Java"],
        ["OSI", "Open Systems Interconnection", "Mô hình tham chiếu kết nối các hệ thống mở 7 tầng"],
        ["API", "Application Programming Interface", "Giao diện lập trình ứng dụng trao đổi dữ liệu RESTful"],
        ["REST", "Representational State Transfer", "Kiến trúc thiết kế dịch vụ web truyền dữ liệu phi trạng thái"],
        ["UUID", "Universally Unique Identifier", "Mã định danh duy nhất toàn cầu 128-bit cho gói tin mạng"],
        ["LAN", "Local Area Network", "Mạng máy tính cục bộ trong văn phòng hoặc phòng lab"],
        ["WAN", "Wide Area Network", "Mạng diện rộng kết nối các mạng máy tính qua khoảng cách lớn"],
        ["CLI", "Command Line Interface", "Giao diện dòng lệnh tương tác trên Terminal/Console"],
        ["UI / UX", "User Interface / User Experience", "Giao diện người dùng và trải nghiệm người dùng"],
        ["ACK", "Acknowledgment", "Gói tin phản hồi báo nhận trong cơ chế bắt tay TCP"],
        ["FIFO", "First In First Out", "Nguyên lý xếp hàng vào trước ra trước trong xử lý luồng"],
        ["RAM", "Random Access Memory", "Bộ nhớ truy xuất ngẫu nhiên tạm thời của máy tính"],
        ["CPU", "Central Processing Unit", "Bộ vi xử lý trung tâm thực thi các phép tính toán băm PoW"]
    ]
    add_table(doc, headers, data, [25, 60, 75])

def add_tables_and_figures_list(doc):
    doc.add_page_break()
    add_heading_1(doc, "DANH MỤC BẢNG BIỂU VÀ HÌNH VẼ")
    
    add_heading_2(doc, "Danh mục Bảng biểu")
    tables_list = [
        ("Bảng 1.1", "Bảng ma trận phân công trách nhiệm nhóm theo chuẩn RACI và GitHub Issues"),
        ("Bảng 1.2", "So sánh đặc tính kỹ thuật giữa TCP Byte Stream và UDP Datagram"),
        ("Bảng 2.1", "Đặc tả các loại thông điệp NetworkMessage trao đổi qua mạng P2P Sockets"),
        ("Bảng 2.2", "Đặc tả các trường thông tin trong cấu trúc khối Block 5 thuộc tính chuẩn Bitcoin"),
        ("Bảng 2.3", "Danh mục các Endpoint REST API của máy chủ NodeApiServer trên cổng 8000"),
        ("Bảng 4.1", "Bảng thống kê các bước kiểm toán của động cơ BlockchainValidator"),
        ("Bảng 4.2", "Kịch bản và kết quả mô phỏng 3 pha tấn công sửa đổi dữ liệu sổ cái"),
        ("Bảng 5.1", "Thông số phần cứng và môi trường phần mềm thử nghiệm hệ thống"),
        ("Bảng 5.2", "Kết quả thực nghiệm 8 bộ kiểm thử tự động toàn diện JUnit 5"),
        ("Bảng 5.3", "So sánh hiệu năng đào PoW giữa Single Thread, Platform Pool và Virtual Threads"),
        ("Bảng 5.4", "Bảng số liệu điều chỉnh độ khó động qua 20 khối liên tiếp (Benchmark Runner)"),
        ("Bảng 5.5", "Bảng đối sánh tính năng giữa hệ thống Đồ án với các giải pháp E-Voting khác")
    ]
    for code, desc in tables_list:
        add_bullet(doc, desc, bold_prefix=code + ": ")

    add_heading_2(doc, "Danh mục Hình vẽ và Sơ đồ")
    figures_list = [
        ("Hình 1.1", "Mô hình phân tầng OSI và vị trí của tầng mạng TCP Sockets và Blockchain"),
        ("Hình 1.2", "Hiện tượng dính gói tin (TCP Sticky Packets) và giải pháp Length-Prefixed Framing"),
        ("Hình 1.3", "Cấu trúc chuỗi khối liên kết băm tiền nhiệm và hiệu ứng sụp đổ dây chuyền"),
        ("Hình 1.4", "Cấu trúc cây nhị phân Merkle Tree và quy trình tính Merkle Root chuẩn Bitcoin"),
        ("Hình 2.1", "Sơ đồ kiến trúc tổng thể 5 tầng của hệ thống E-Voting Blockchain P2P"),
        ("Hình 2.2", "Biểu đồ tuần tự (Sequence Diagram) quy trình cử tri bỏ phiếu đến khi đóng khối"),
        ("Hình 2.3", "Biểu đồ tuần tự quá trình bắt tay và đồng bộ chuỗi dài nhất giữa các nút P2P"),
        ("Hình 3.1", "Kiến trúc xử lý socket bất đồng bộ trên Java 21 Virtual Threads của NodeServer"),
        ("Hình 3.2", "Cơ chế truyền tin đồn Gossip Protocol và bộ đệm khử lặp seenMessageIds"),
        ("Hình 4.1", "Thuật toán phân chia không gian Nonce theo strides 4096 của MultiThreadedMiner"),
        ("Hình 4.2", "Đồ thị điều chỉnh độ khó động theo chu kỳ Epoch và kẹp hệ số giảm chấn"),
        ("Hình 5.1", "Giao diện Bỏ phiếu trực tuyến dành cho Cử tri (Ballot Portal - /vote)"),
        ("Hình 5.2", "Giao diện Sơ đồ cụm mạng 10 nút P2P Mesh trực quan hóa hạt photon (/nodes)"),
        ("Hình 5.3", "Giao diện Sổ cái chuỗi khối Explorer kiểm toán dữ liệu thời gian thực (/)")
    ]
    for code, desc in figures_list:
        add_bullet(doc, desc, bold_prefix=code + ": ")

def add_table_of_contents(doc, page_map=None):
    doc.add_page_break()
    add_heading_1(doc, "MỤC LỤC TỔNG QUAN BÁO CÁO")
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r = p_sub.add_run("(Hệ thống đề mục chi tiết toàn bộ 6 chương nội dung chuyên sâu)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = COLOR_MUTED

    # Default fallback page map if not provided
    pm = page_map or {
        "EVAL": "iii", "DECL": "iv", "THANKS": "v", "ABS": "vi", "ABBR": "viii", "FIG": "ix",
        "C1": "1", "C1_1": "1", "C1_2": "2", "C1_3": "3", "C1_4": "4", "C1_5": "6", "C1_6": "7", "C1_7": "8",
        "C2": "10", "C2_1": "10", "C2_2": "11", "C2_3": "12", "C2_4": "13", "C2_5": "15", "C2_6": "16", "C2_7": "17", "C2_8": "17",
        "C3": "19", "C3_1": "19", "C3_2": "20", "C3_3": "21", "C3_4": "22", "C3_5": "23", "C3_6": "24", "C3_7": "25",
        "C4": "26", "C4_1": "26", "C4_2": "27", "C4_3": "28", "C4_4": "29", "C4_5": "30", "C4_6": "31", "C4_7": "32",
        "C5": "33", "C5_1": "33", "C5_2": "34", "C5_3": "36", "C5_4": "38",
        "C6": "39", "C6_1": "39", "C6_2": "40", "C6_3": "40",
        "REF": "41", "APP": "42"
    }

    toc_items = [
        ("PHIẾU ĐÁNH GIÁ KẾT QUẢ ĐỒ ÁN", pm.get("EVAL", "iii"), 0),
        ("LỜI CAM ĐOAN", pm.get("DECL", "iv"), 0),
        ("LỜI CẢM ƠN", pm.get("THANKS", "v"), 0),
        ("TÓM TẮT ĐỀ TÀI (ABSTRACT)", pm.get("ABS", "vi"), 0),
        ("DANH MỤC CÁC TỪ VIẾT TẮT", pm.get("ABBR", "viii"), 0),
        ("DANH MỤC BẢNG BIỂU VÀ HÌNH VẼ", pm.get("FIG", "ix"), 0),
        ("CHƯƠNG 1: TỔNG QUAN VỀ ĐỀ TÀI VÀ CƠ SỞ LÝ THUYẾT", pm.get("C1", "1"), 0),
        ("1.1. Bối cảnh và tính cấp thiết của đề tài", pm.get("C1_1", "1"), 1),
        ("1.2. Mục tiêu nghiên cứu và giới hạn phạm vi đề tài", pm.get("C1_2", "2"), 1),
        ("1.3. Phân công trách nhiệm và đóng góp của các thành viên trong nhóm", pm.get("C1_3", "3"), 1),
        ("1.4. Cơ sở lý thuyết về Mạng máy tính và Giao thức TCP/IP", pm.get("C1_4", "4"), 1),
        ("1.5. Cơ sở lý thuyết về Công nghệ Blockchain và Mật mã học", pm.get("C1_5", "6"), 1),
        ("1.6. Cơ chế đồng thuận Proof-of-Work (PoW) và Độ khó động", pm.get("C1_6", "7"), 1),
        ("1.7. Nền tảng công nghệ Java 21 hiện đại và Virtual Threads", pm.get("C1_7", "8"), 1),
        ("CHƯƠNG 2: PHÂN TÍCH YÊU CẦU VÀ THIẾT KẾ KIẾN TRÚC HỆ THỐNG", pm.get("C2", "10"), 0),
        ("2.1. Phân tích yêu cầu chức năng (Functional Requirements)", pm.get("C2_1", "10"), 1),
        ("2.2. Phân tích yêu cầu phi chức năng (Non-Functional Requirements)", pm.get("C2_2", "11"), 1),
        ("2.3. Thiết kế kiến trúc tổng thể 5 tầng của hệ thống", pm.get("C2_3", "12"), 1),
        ("2.4. Thiết kế chi tiết Tầng mạng P2P TCP Sockets", pm.get("C2_4", "13"), 1),
        ("2.5. Thiết kế chi tiết Tầng lõi Dữ liệu và Mật mã học", pm.get("C2_5", "15"), 1),
        ("2.6. Thiết kế chi tiết Tầng Đồng thuận PoW và Điều chỉnh độ khó", pm.get("C2_6", "16"), 1),
        ("2.7. Thiết kế Tầng Kiểm toán và Giả lập Tấn công", pm.get("C2_7", "17"), 1),
        ("2.8. Thiết kế Tầng Lưu trữ bền vững và Web REST API", pm.get("C2_8", "17"), 1),
        ("CHƯƠNG 3: HIỆN THỰC HÓA CHI TIẾT TẦNG MẠNG P2P TCP SOCKETS", pm.get("C3", "19"), 0),
        ("3.1. Hiện thực hóa Giao thức đóng khung luồng tin NetworkProtocol", pm.get("C3_1", "19"), 1),
        ("3.2. Hiện thực hóa Khung thông điệp NetworkMessage", pm.get("C3_2", "20"), 1),
        ("3.3. Hiện thực hóa Máy chủ TCP bất đồng bộ NodeServer trên Java 21", pm.get("C3_3", "21"), 1),
        ("3.4. Hiện thực hóa Kênh kết nối song công PeerConnection", pm.get("C3_4", "22"), 1),
        ("3.5. Hiện thực hóa Nút mạng P2P trung tâm và Giao thức Gossip", pm.get("C3_5", "23"), 1),
        ("3.6. Hiện thực hóa Hòm phiếu đệm Mempool an toàn luồng", pm.get("C3_6", "24"), 1),
        ("3.7. Hiện thực hóa Nút chạy độc lập đa máy tính StandaloneNode", pm.get("C3_7", "25"), 1),
        ("CHƯƠNG 4: HIỆN THỰC HÓA CÁC MODULE LÕI, ĐỒNG THUẬN VÀ BẢO MẬT", pm.get("C4", "26"), 0),
        ("4.1. Hiện thực hóa Cấu trúc Khối Block và Sổ cái Bầu cử Blockchain", pm.get("C4_1", "26"), 1),
        ("4.2. Hiện thực hóa Mật mã học hiệu năng cao và Cây Merkle", pm.get("C4_2", "27"), 1),
        ("4.3. Hiện thực hóa Bộ đào PoW đa luồng MultiThreadedMiner", pm.get("C4_3", "28"), 1),
        ("4.4. Hiện thực hóa Động cơ điều chỉnh độ khó động DifficultyAdjuster", pm.get("C4_4", "29"), 1),
        ("4.5. Hiện thực hóa Bộ kiểm toán Sổ cái toàn diện BlockchainValidator", pm.get("C4_5", "30"), 1),
        ("4.6. Hiện thực hóa Bộ mô phỏng tấn công giả mạo AttackSimulator", pm.get("C4_6", "31"), 1),
        ("4.7. Hiện thực hóa Lưu trữ bền vững nguyên tử LedgerStorage", pm.get("C4_7", "32"), 1),
        ("CHƯƠNG 5: HIỆN THỰC GIAO DIỆN NGƯỜI DÙNG VÀ KẾT QUẢ THỰC NGHIỆM", pm.get("C5", "33"), 0),
        ("5.1. Hiện thực hóa Máy chủ Web API NodeApiServer", pm.get("C5_1", "33"), 1),
        ("5.2. Các giao diện trực quan hóa hệ sinh thái (Vote, Nodes, Explorer)", pm.get("C5_2", "34"), 1),
        ("5.3. Môi trường và các kịch bản thực nghiệm chuyên sâu", pm.get("C5_3", "36"), 1),
        ("5.4. Đánh giá và so sánh kết quả thực nghiệm", pm.get("C5_4", "38"), 1),
        ("CHƯƠNG 6: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", pm.get("C6", "39"), 0),
        ("6.1. Các kết quả nổi bật đạt được của đề tài", pm.get("C6_1", "39"), 1),
        ("6.2. Các hạn chế còn tồn tại và bài học kinh nghiệm", pm.get("C6_2", "40"), 1),
        ("6.3. Hướng phát triển trong tương lai", pm.get("C6_3", "40"), 1),
        ("TÀI LIỆU THAM KHẢO", pm.get("REF", "41"), 0),
        ("PHỤ LỤC: HƯỚNG DẪN CÀI ĐẶT VÀ VẬN HÀNH 1-CLICK", pm.get("APP", "42"), 0)
    ]
    
    for title, page_str, level in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0.5 if level > 0 else 1.5)
        p.paragraph_format.space_after = Pt(0.5)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Mm(level * 4)
        p.paragraph_format.first_line_indent = Mm(0)
        
        r_tit = p.add_run(title)
        r_tit.font.name = "Times New Roman"
        r_tit.font.size = Pt(9.2 if level == 0 else 8.8)
        r_tit.font.bold = (level == 0)
        r_tit.font.color.rgb = COLOR_PRIMARY if level == 0 else COLOR_DARK
        
        dot_count = max(4, 82 - len(title) - level * 4)
        r_tab = p.add_run("  " + ("." * dot_count) + "  " + page_str)
        r_tab.font.name = "Times New Roman"
        r_tab.font.size = Pt(8.8)
        r_tab.font.color.rgb = COLOR_MUTED
