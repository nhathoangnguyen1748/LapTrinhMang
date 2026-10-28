# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

COLOR_PRIMARY = RGBColor(0, 32, 96)       # Navy Blue #002060
COLOR_SECONDARY = RGBColor(31, 78, 121)   # Slate Blue #1F4E79
COLOR_ACCENT = RGBColor(46, 117, 182)     # Steel Blue #2E75B6
COLOR_DARK = RGBColor(38, 38, 38)         # #262626
COLOR_MUTED = RGBColor(89, 89, 89)        # #595959

def setup_initial_document(doc):
    # Cấu hình Section 0 (Trang bìa chính)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""

def setup_body_section(doc):
    # Tạo Section mới bắt đầu từ Trang 3 trở đi cho toàn bộ nội dung báo cáo
    section = doc.add_section()
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)
    
    # QUAN TRỌNG: Xóa triệt để pgBorders để toàn bộ các trang nội dung KHÔNG có khung viền bìa!
    pgB = section._sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgBorders')
    if pgB is not None:
        section._sectPr.remove(pgB)
        
    # Thiết lập chân trang (Footer) độc lập, đánh số trang: Trang X / Y
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    
    run_l = p.add_run("Báo cáo Đồ án Lập trình mạng - E-Voting Blockchain P2P  |  Trang ")
    run_l.font.name = "Times New Roman"
    run_l.font.size = Pt(9.5)
    run_l.font.italic = True
    run_l.font.color.rgb = COLOR_MUTED
    
    fld_page = parse_xml(r'<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="PAGE"/>')
    p._p.append(fld_page)
    
    run_m = p.add_run(" / ")
    run_m.font.name = "Times New Roman"
    run_m.font.size = Pt(9.5)
    run_m.font.italic = True
    run_m.font.color.rgb = COLOR_MUTED
    
    fld_numpages = parse_xml(r'<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="NUMPAGES"/>')
    p._p.append(fld_numpages)
    
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:top w:val="single" w:sz="4" w:space="3" w:color="D3D3D3"/>'
                     r'</w:pBdr>')
    p._p.get_or_add_pPr().append(pBdr)
    return section

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text.upper())
    run.font.name = "Times New Roman"
    run.font.size = Pt(14.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = COLOR_ACCENT
    return p

def add_heading_4(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK
    return p

def add_paragraph(doc, text, bold_prefix=None, italic=False, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Mm(8)
    else:
        p.paragraph_format.first_line_indent = Mm(0)
        
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
        
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.italic = italic
    run.font.color.rgb = COLOR_DARK
    return p

def add_bullet(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Mm(9 + level * 5)
    p.paragraph_format.first_line_indent = Mm(-5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    bullet_symbol = "• " if level == 0 else "– "
    r_sym = p.add_run(bullet_symbol)
    r_sym.font.name = "Times New Roman"
    r_sym.font.size = Pt(11.5)
    r_sym.font.bold = True
    r_sym.font.color.rgb = COLOR_ACCENT
    
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(11.5)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_DARK
        
    r_t = p.add_run(text)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11.5)
    r_t.font.color.rgb = COLOR_DARK
    return p

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(5)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.left_indent = Mm(8)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_SECONDARY
        
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Mm(160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="F4F6F9"/>'.format(nsdecls('w'))))
    tcBorders = parse_xml(r'<w:tcBorders {} >'
                          r'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
                          r'<w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E79"/>'
                          r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
                          r'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
                          r'</w:tcBorders>'.format(nsdecls('w')))
    tcPr.append(tcBorders)
    
    tcMar = parse_xml(r'<w:tcMar {} >'
                      r'<w:top w:w="60" w:type="dxa"/>'
                      r'<w:left w:w="120" w:type="dxa"/>'
                      r'<w:bottom w:w="60" w:type="dxa"/>'
                      r'<w:right w:w="120" w:type="dxa"/>'
                      r'</w:tcMar>'.format(nsdecls('w')))
    tcPr.append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.first_line_indent = Mm(0)
    
    lines = code_str.strip().split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.first_line_indent = Mm(0)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(36, 41, 47)
        
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(3)

def add_table(doc, headers, data, col_widths=None, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_SECONDARY

    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Mm(0)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls('w'))))
        
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_color = "F2F5F9" if (row_idx % 2 == 1) else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 or len(str(cell_value)) < 15 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.line_spacing = 1.08
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.0)
                r.font.color.rgb = COLOR_DARK
                
            tcPr = row_cells[col_idx]._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_color)))

    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(r'<w:tcBorders {} >'
                                r'<w:top w:val="single" w:sz="4" w:space="0" w:color="B0C4DE"/>'
                                r'<w:left w:val="single" w:sz="4" w:space="0" w:color="B0C4DE"/>'
                                r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="B0C4DE"/>'
                                r'<w:right w:val="single" w:sz="4" w:space="0" w:color="B0C4DE"/>'
                                r'</w:tcBorders>'.format(nsdecls('w')))
            tcPr.append(borders)
            
            tcMar = parse_xml(r'<w:tcMar {} >'
                              r'<w:top w:w="40" w:type="dxa"/>'
                              r'<w:left w:w="70" w:type="dxa"/>'
                              r'<w:bottom w:w="40" w:type="dxa"/>'
                              r'<w:right w:w="70" w:type="dxa"/>'
                              r'</w:tcMar>'.format(nsdecls('w')))
            tcPr.append(tcMar)
            
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Mm(w)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(3)

def add_callout(doc, text, title="LƯU Ý THIẾT KẾ QUAN TRỌNG"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Mm(160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="EBF3FB"/>'.format(nsdecls('w'))))
    tcBorders = parse_xml(r'<w:tcBorders {} >'
                          r'<w:top w:val="none"/>'
                          r'<w:left w:val="single" w:sz="36" w:space="0" w:color="2E75B6"/>'
                          r'<w:bottom w:val="none"/>'
                          r'<w:right w:val="none"/>'
                          r'</w:tcBorders>'.format(nsdecls('w')))
    tcPr.append(tcBorders)
    
    tcMar = parse_xml(r'<w:tcMar {} >'
                      r'<w:top w:w="60" w:type="dxa"/>'
                      r'<w:left w:w="120" w:type="dxa"/>'
                      r'<w:bottom w:w="60" w:type="dxa"/>'
                      r'<w:right w:w="120" w:type="dxa"/>'
                      r'</w:tcMar>'.format(nsdecls('w')))
    tcPr.append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.first_line_indent = Mm(0)
    r_tit = p.add_run("📌 " + title + "\n")
    r_tit.font.name = "Times New Roman"
    r_tit.font.size = Pt(10.5)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY
    
    r_txt = p.add_run(text)
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = COLOR_DARK
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(3)

def add_figure_box(doc, ascii_art, caption):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Mm(160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd {} w:fill="FAFBFD"/>'.format(nsdecls('w'))))
    tcBorders = parse_xml(r'<w:tcBorders {} >'
                          r'<w:top w:val="single" w:sz="6" w:space="0" w:color="2E75B6"/>'
                          r'<w:left w:val="single" w:sz="6" w:space="0" w:color="2E75B6"/>'
                          r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="2E75B6"/>'
                          r'<w:right w:val="single" w:sz="6" w:space="0" w:color="2E75B6"/>'
                          r'</w:tcBorders>'.format(nsdecls('w')))
    tcPr.append(tcBorders)
    
    tcMar = parse_xml(r'<w:tcMar {} >'
                      r'<w:top w:w="60" w:type="dxa"/>'
                      r'<w:left w:w="100" w:type="dxa"/>'
                      r'<w:bottom w:w="60" w:type="dxa"/>'
                      r'<w:right w:w="100" w:type="dxa"/>'
                      r'</w:tcMar>'.format(nsdecls('w')))
    tcPr.append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    lines = ascii_art.strip().split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Mm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(7.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(16, 44, 87)
        
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(3)
    p_cap.paragraph_format.space_after = Pt(5)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_SECONDARY
