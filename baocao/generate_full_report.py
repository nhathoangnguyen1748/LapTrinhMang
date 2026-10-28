# -*- coding: utf-8 -*-
import os
import sys
import docx
from docx.shared import Mm, Pt

from report_styles import setup_initial_document, setup_body_section
from make_preamble import (
    add_cover_page, add_eval_sheet, add_declarations,
    add_acknowledgements, add_abstract, add_abbreviations_table,
    add_tables_and_figures_list, add_table_of_contents
)
from chapter1 import add_chapter_1
from chapter2 import add_chapter_2
from chapter3 import add_chapter_3
from chapter4 import add_chapter_4
from chapter5 import add_chapter_5
from chapter6 import add_chapter_6

def build_report():
    print("=" * 70)
    print(" DANG XAY DUNG BAO CAO DO AN LAP TRINH MANG BLOCKCHAIN P2P (JAVA 21) ")
    print("=======================================================================")

    doc = docx.Document()
    
    # 1. Cấu hình Section 0 (Trang bìa chính)
    setup_initial_document(doc)
    
    logo_path = "/sessions/friendly-kind-goodall/mnt/LapTrinhMang/baocao/hutech_logo.png"
    
    # Bảng phân trang đối soát chính xác 100%
    page_map = {
        "EVAL": "iii", "DECL": "iv", "THANKS": "v", "ABS": "vi", "ABBR": "viii", "FIG": "ix",
        "C1": "11", "C1_1": "11", "C1_2": "11", "C1_3": "12", "C1_4": "13", "C1_5": "15", "C1_6": "16", "C1_7": "17",
        "C2": "19", "C2_1": "19", "C2_2": "19", "C2_3": "20", "C2_4": "21", "C2_5": "22", "C2_6": "22", "C2_7": "23", "C2_8": "23",
        "C3": "25", "C3_1": "25", "C3_2": "26", "C3_3": "27", "C3_4": "28", "C3_5": "29", "C3_6": "31", "C3_7": "31",
        "C4": "32", "C4_1": "32", "C4_2": "33", "C4_3": "35", "C4_4": "36", "C4_5": "36", "C4_6": "37", "C4_7": "37",
        "C5": "38", "C5_1": "38", "C5_2": "38", "C5_3": "40", "C5_4": "42",
        "C6": "43", "C6_1": "43", "C6_2": "43", "C6_3": "44",
        "REF": "45", "APP": "46"
    }

    print(" [1/9] Tao Trang bia chinh HUTECH (khung vien hoa van kep)...")
    add_cover_page(doc, logo_path, is_subcover=False)
    
    print(" [2/9] Tao Trang bia phu (Bia lot ben trong co khung vien)...")
    add_cover_page(doc, logo_path, is_subcover=True)
    
    print(" [*] Tao Section moi cho toan bo noi dung (Trang 3 tro di - KHONG KHUNG VIEN)...")
    setup_body_section(doc)
    
    print(" [3/9] Tao Phieu danh gia GVHD, Loi cam doan, Loi cam on, Abstract...")
    add_eval_sheet(doc)
    add_declarations(doc)
    add_acknowledgements(doc)
    add_abstract(doc)
    add_abbreviations_table(doc)
    add_tables_and_figures_list(doc)
    add_table_of_contents(doc, page_map)
    
    print(" [4/9] Tao Chuong 1: Tong quan de tai va Co so ly thuyet...")
    add_chapter_1(doc)
    
    print(" [5/9] Tao Chuong 2: Phan tich yeu cau va Thiet ke kien truc...")
    add_chapter_2(doc)
    
    print(" [6/9] Tao Chuong 3: Hien thuc hoa chi tiet Tang mang P2P TCP Sockets...")
    add_chapter_3(doc)
    
    print(" [7/9] Tao Chuong 4: Hien thuc hoa cac Module loi, Dong thuan va Bao mat...")
    add_chapter_4(doc)
    
    print(" [8/9] Tao Chuong 5: Hien thuc Giao dien nguoi dung va Ket qua thuc nghiem...")
    add_chapter_5(doc)
    
    print(" [9/9] Tao Chuong 6: Ket luan, Huong phat trien, Tai lieu tham khao & Phu luc...")
    add_chapter_6(doc)
    
    output_docx = "/sessions/friendly-kind-goodall/mnt/LapTrinhMang/baocao/BaoCao_DoAn_LapTrinhMang_Blockchain_EVoting.docx"
    doc.save(output_docx)
    print("\n" + "=" * 70)
    print(f" [THANH CONG] Da xuat file bao cao DOCX tai:\n -> {output_docx}")
    print("=" * 70)

if __name__ == "__main__":
    build_report()
